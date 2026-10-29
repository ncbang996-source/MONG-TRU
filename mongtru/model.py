"""Thông số móng trụ rời (4 trụ): giá trị mặc định, đọc Excel (sheet TONG HOP), kiểm tra, khối lượng.

Mỗi chân cột có 1 trụ móng: bản đế vuông B x B (mặt trên vát từ hb ở gờ quanh cổ xuống hv ở mép), cổ móng
b1 (đỉnh) / b2 (chân). 4 trụ đặt tại 4 góc hình vuông cạnh Bt (khoảng cách tim trụ).
Khối lượng bê tông, lót, đào, đắp, ván khuôn theo công thức sheet "KL" của Excel mẫu."""
from __future__ import annotations

import math
import re

# ô Excel -> thông số (sheet TONG HOP)
CELLS_NUM = {
    "B": "B9",      # Bm = Lm: cạnh bản đế
    "H1": "B10",    # Hm: chiều sâu chôn móng
    "H2": "B11",    # d: cao cổ móng trên mặt đất
    "BT": "B12",    # Bt: khoảng cách tim trụ
    "B1": "B13",    # b1: cạnh cổ móng ở đỉnh
    "B2": "B14",    # b2: cạnh cổ móng ở chân
    "HB": "B15",    # hb: cao bản đế tại gờ cổ móng
    "HV": "B16",    # hv: cao bản đế ở mép
    "SLOPE": "C18", "MPIT": "C19",
}
CELLS_STR = {"GD": "C6", "COT": "C4", "TENMONG": "C5", "BTM": "J4", "BTL": "J5", "CODE": "J6"}
CELL_HT = "B4"
CELLS_MABV = ("J7", "K7")                          # mã số bản vẽ (khung tên) = J7 & K7
CELLS_BLN = ("C20", "D20")
# cốt thép: (khoá, ô) - Ø, bước (mm), số thanh
CELLS_REBAR = [
    ("RB_BOT_D", "H21"), ("RB_BOT_S", "I21"), ("RB_TOP_D", "H22"), ("RB_TOP_S", "I22"),
    ("RB_CH_D", "H23"), ("RB_CH_S", "I23"),
    ("RC_D", "H25"), ("RC_N", "I25"), ("RT_D", "H26"), ("RT_S", "I26"),
    ("RP_D", "H27"), ("RP_S", "I27"), ("RL_D", "H28"), ("RL_S", "I28"),
]
CELLS_COVER = {"COVER_B": "C22", "COVER": "C23"}   # mm trong Excel -> m
CELL_RLAYERS = "K28"                               # "02 lớp"
CELL_NPIER = "C24"                                 # số trụ móng
CELL_DP = "H29"                                    # dự phòng thép (%)

HT_NAMES = ["ĐỠ", "NÉO GÓC", "NÉO DỪNG"]
SCALES = [50, 75, 100, 125, 150, 200, 250]
SLOPE_DRAW = 0.5          # hệ số mở mái THỂ HIỆN trên bản vẽ (C18 chỉ dùng tính khối lượng)
TL = 0.1                  # dày bê tông lót (m)
ML = 0.1                  # lót tràn ra mỗi bên (m)
MPL = 0.1                 # gờ phẳng quanh chân cổ móng trên mặt bản đế (m)


def default_params() -> dict:
    return {
        "B": 2.7, "H1": 2.5, "H2": 0.5, "BT": 3.2, "B1": 0.7, "B2": 0.7, "HB": 0.7, "HV": 0.3,
        "SLOPE": 0.5, "MPIT": 0.3, "NPIER": 4.0, "DP": 10.0,
        "HT": 0, "BLN": "BLN-M36-250", "GD": "TKBVTC", "TENMONG": "",
        "SCALE": 100,
        "BTM": "B20", "BTL": "B7.5",
        "COT": "D112-30", "CODE": "",
        "RB_BOT_D": 12.0, "RB_BOT_S": 200.0, "RB_TOP_D": 12.0, "RB_TOP_S": 200.0,
        "RB_CH_D": 12.0, "RB_CH_S": 600.0,
        "RC_D": 20.0, "RC_N": 20.0, "RT_D": 8.0, "RT_S": 200.0, "RP_D": 8.0, "RP_S": 400.0,
        "RL_D": 8.0, "RL_S": 100.0, "RL_N": 2.0,
        "COVER": 0.05, "COVER_B": 0.05,
        # khung tên
        "MABV": "",
        "PIT_DRAW": 0.0,    # trình bày hố đào trên hình: 0 = tự động (theo cách tính), 1 = đào chung, 2 = đào riêng
    }


# thông số không có trong Excel: giữ lại khi mở file Excel khác
KEEP_KEYS = ("MABV", "PIT_DRAW")
PIT_DRAW_NAMES = ["Tự động", "Đào chung 1 hố", "Đào riêng từng trụ"]


def fresh_params(prev: dict | None = None) -> dict:
    """Nền khi mở file Excel mới: thông số mặc định + các thông số không có trong Excel của lần trước."""
    p = default_params()
    if prev:
        p.update({k: prev[k] for k in KEEP_KEYS if k in prev})
    return p


def type_name(p: dict) -> str:
    if p.get("TENMONG"):
        return p["TENMONG"]
    return f"MT{p['BT'] * 10:.1f} 4T{p['B'] * 10:.0f}-{H(p) * 10:.0f}"


def ht_from_string(s: str) -> int:
    u = s.upper()
    if "DỪNG" in u or "DUNG" in u:
        return 2
    if "NÉO" in u or "NEO" in u or "GÓC" in u or "GOC" in u:
        return 1
    return 0


def _num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _open_sheet(path: str, sheet: str):
    """Trả về hàm đọc ô theo địa chỉ ("B9") của sheet; hỗ trợ .xlsx/.xlsm (openpyxl) và .xls (xlrd)."""
    if path.lower().endswith(".xls"):
        import xlrd
        wb = xlrd.open_workbook(path)
        if sheet not in wb.sheet_names():
            raise ValueError(f'Không tìm thấy sheet "{sheet}" trong file Excel.')
        ws = wb.sheet_by_name(sheet)

        def val(a):
            m = re.fullmatch(r"([A-Z]+)(\d+)", a)
            col = 0
            for ch in m.group(1):
                col = col * 26 + ord(ch) - 64
            r, c = int(m.group(2)) - 1, col - 1
            if r >= ws.nrows or c >= ws.ncols:
                return None
            cell = ws.cell(r, c)
            if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                return None
            return cell.value
        return val, lambda: None
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if sheet not in wb.sheetnames:
        raise ValueError(f'Không tìm thấy sheet "{sheet}" trong file Excel.')
    ws = wb[sheet]
    return (lambda a: ws[a].value), wb.close


def read_excel(path: str, sheet: str = "TONG HOP", p: dict | None = None) -> tuple[dict, list[str]]:
    """Đọc sheet TONG HOP. Trả về (thông số, danh sách cảnh báo)."""
    p = dict(p or default_params())
    warn: list[str] = []
    val, close = _open_sheet(path, sheet)

    nread = 0
    for k, a in CELLS_NUM.items():
        v = val(a)
        if _num(v):
            p[k] = float(v)
            nread += 1
        else:
            warn.append(f"Ô {a} ({k}) không có số, giữ giá trị {p[k]:g}.")
    for k, a in CELLS_STR.items():
        v = val(a)
        if v not in (None, ""):
            p[k] = str(v).strip()
    mv = [val(a) for a in CELLS_MABV]
    if mv[1] not in (None, ""):
        num = f"{mv[1]:g}" if _num(mv[1]) else str(mv[1]).strip()
        p["MABV"] = f"{str(mv[0] or '').strip()}{num}"
    v = val(CELL_HT)
    if isinstance(v, str):
        p["HT"] = ht_from_string(v)
    bln = "".join(str(val(a) or "").strip() for a in CELLS_BLN)
    if bln:
        p["BLN"] = bln
    for k, a in CELLS_REBAR:
        v = val(a)
        if _num(v) and v > 0:
            p[k] = float(v)
        else:
            warn.append(f"Ô {a} ({k}) không có số, giữ giá trị {p[k]:g}.")
    for k, a in CELLS_COVER.items():
        v = val(a)
        if _num(v) and v > 0:
            p[k] = float(v) / 1000.0
        else:
            warn.append(f"Ô {a} (lớp bảo vệ) không có số, giữ giá trị {p[k] * 1000:g} mm.")
    v = val(CELL_RLAYERS)
    if v not in (None, ""):
        try:
            n = int(v) if _num(v) else int("".join(c for c in str(v) if c.isdigit()) or 0)
            if n > 0:
                p["RL_N"] = float(n)
        except ValueError:
            pass
    v = val(CELL_NPIER)
    if _num(v) and v > 0:
        p["NPIER"] = float(v)
    else:
        warn.append(f"Ô {CELL_NPIER} (số trụ móng) không có số, giữ giá trị {p['NPIER']:g}.")
    v = val(CELL_DP)
    if _num(v) and v >= 0:
        p["DP"] = float(v)
    else:
        warn.append(f"Ô {CELL_DP} (dự phòng thép) không có số, giữ giá trị {p['DP']:g}%.")
    close()
    if nread == 0:
        raise ValueError("Không đọc được số liệu nào từ sheet. Kiểm tra lại địa chỉ ô.")
    return p, warn


def H(p):
    return p["H1"] + p["H2"]


def npier(p) -> int:
    return int(round(p.get("NPIER", 4)))


# ---------------------------------------------------------------------------
# khối lượng (sheet "KL"), cho cả móng (NPIER trụ)
def _frustum_sq(a, b, h):
    """Thể tích chóp cụt đáy vuông cạnh a, b, cao h."""
    return h * (a * a + b * b + a * b) / 3.0


def vol_pad(p) -> float:
    """Bản đế 1 trụ: phần đứng cao hv + phần vát (chóp cụt từ B lên gờ b2 + 2·mpl)."""
    b, a = p["B"], p["B2"] + 2 * MPL
    return b * b * p["HV"] + _frustum_sq(b, a, p["HB"] - p["HV"])


def vol_col(p) -> float:
    """Cổ móng 1 trụ (chóp cụt b2 -> b1) từ mặt gờ hb lên đỉnh."""
    return _frustum_sq(p["B2"], p["B1"], H(p) - p["HB"])


def b_at(p, z):
    """Cạnh cổ móng tại cao độ z (tính từ đáy bản đế)."""
    t = (z - p["HB"]) / max(1e-9, H(p) - p["HB"])
    return p["B2"] + (p["B1"] - p["B2"]) * min(1.0, max(0.0, t))


def vol_mong(p) -> float:
    return npier(p) * (vol_pad(p) + vol_col(p))


def vol_above(p) -> float:
    """Thể tích cổ móng trên mặt đất (cả móng)."""
    return npier(p) * _frustum_sq(b_at(p, p["H1"]), p["B1"], p["H2"])


def vol_lot(p) -> float:
    return npier(p) * (p["B"] + 2 * ML) ** 2 * TL


def pit_depth(p) -> float:
    return p["H1"] + TL


def pad_dist(p) -> float:
    """Khoảng cách tim bản đế (m) = Bt + (b2 − b1): cổ vát ra phía ngoài, mặt trong thẳng đứng."""
    return p["BT"] + p["B2"] - p["B1"]


def pit_joint(p) -> bool:
    """True: đào chung 1 hố cho cả 4 trụ (miệng các hố riêng chạm nhau), False: đào từng trụ."""
    gap = pad_dist(p) - 2 * (0.5 * p["B"] + p["MPIT"] + pit_depth(p) * p["SLOPE"])
    return gap <= 0


def pit_bottom(p) -> float:
    """Cạnh đáy hố đào (m): hố chung hoặc hố từng trụ."""
    base = p["B"] + 2 * p["MPIT"]
    return base + pad_dist(p) if pit_joint(p) else base


def _frustum_rect(h, u, v, w, x):
    return h / 6 * (u * v + w * x + (u + w) * (v + x))


def vol_dao(p) -> float:
    a = pit_bottom(p)
    t = a + 2 * pit_depth(p) * p["SLOPE"]
    v1 = _frustum_rect(pit_depth(p), a, a, t, t)
    n = npier(p)
    return v1 * (n / 4.0 if pit_joint(p) else n)


def vol_dap(p) -> float:
    return vol_dao(p) - vol_lot(p) - (vol_mong(p) - vol_above(p))


def area_vk(p) -> float:
    """Ván khuôn: 4 mặt cổ móng theo diện tích thật (2 mặt trong đứng, 2 mặt ngoài vát - hình thang đáy b1, b2)
    + 4 mặt mép bản đế (mặt vát bản đế không ghép ván khuôn)."""
    hc = H(p) - p["HB"]
    b1, b2 = p["B1"], p["B2"]
    side = (b1 + b2) / 2 * (hc + math.hypot(hc, b2 - b1))
    return npier(p) * (2 * side + 4 * p["B"] * p["HV"])


def den_bu(p) -> float:
    """Cạnh đền bù (m): bao 4 trụ + 0.5 m mỗi bên."""
    return pad_dist(p) + p["B"] + 1.0


def bolt_dia(p) -> int:
    m = re.search(r"M\s*(\d+)", str(p.get("BLN", "")))
    return int(m.group(1)) if m else 36


def bolt_gap(p, b) -> float:
    """Khoảng cách giữa 2 bu lông 1 cụm (m): số sau dấu '-' của 'M36-250', không có thì theo cổ móng."""
    m = re.search(r"M\s*\d+\s*-\s*(\d+)", str(p.get("BLN", "")))
    if m:
        return int(m.group(1)) / 1000.0
    return 0.47 * b


def bolt_text(p) -> str:
    return f"{4 * npier(p)}M{bolt_dia(p)}"


# ---------------------------------------------------------------------------
def validate(p) -> str | None:
    keys = ("B", "H1", "H2", "BT", "B1", "B2", "HB", "HV")
    if min(p[k] for k in keys) < 1e-6:
        return "Có kích thước <= 0. Kiểm tra lại các ô nhập."
    if p["B1"] > p["B2"] + 1e-9:
        return "Cạnh cổ móng ở đỉnh b1 không được lớn hơn cạnh ở chân b2."
    if p["SLOPE"] < 1e-6:
        return "Hệ số mở mái phải > 0."
    if p["MPIT"] < ML:
        return "Mở rộng đáy hố đào phải >= lót ra mỗi bên (0.1 m)."
    if p["HV"] > p["HB"]:
        return "hv (cao mép bản đế) không được lớn hơn hb."
    if p["HB"] >= p["H1"]:
        return "Hm (chiều sâu chôn) phải lớn hơn hb."
    if p["B"] < p["B2"] + 2 * MPL + 0.2:
        return "Bm quá nhỏ so với cổ móng."
    if pad_dist(p) < p["B"] + 0.1:
        return "Khoảng cách tim bản đế Bt + (b2 − b1) phải lớn hơn cạnh bản đế Bm (các bản đế chồng nhau)."
    n = npier(p)
    if n < 4 or n % 4:
        return "Số trụ móng phải là bội số của 4."
    keys = ("RB_BOT_D", "RB_BOT_S", "RB_TOP_D", "RB_TOP_S", "RB_CH_D", "RB_CH_S", "RC_D",
            "RT_D", "RT_S", "RP_D", "RP_S", "RL_D", "RL_S")
    if min(p[k] for k in keys) < 1e-6:
        return "Đường kính / khoảng cách thép phải > 0."
    n = int(round(p["RC_N"]))
    if n < 4 or n % 4:
        return "Số thanh thép chủ cổ móng phải là bội số của 4 (>= 4)."
    if p["RL_N"] < 1:
        return "Số lớp lưới thép đầu trụ phải >= 1."
    if p["COVER"] < 0 or p["COVER_B"] < 0:
        return "Lớp BT bảo vệ không được âm."
    if p.get("DP", 0) < 0:
        return "Dự phòng thép không được âm."
    return None


GAP_MIN = 0.04      # khoảng hở ngang tối thiểu giữa các thanh thép chủ (m)


def warnings(p) -> list:
    """Cảnh báo (không chặn dựng bản vẽ)."""
    out = []
    try:
        m = 1 + int(round(p["RC_N"])) // 4
        d = p["RC_D"] / 1000.0
        xi = p["B1"] / 2 - p["COVER"] - d / 2
        if m > 1:
            gap = 2 * xi / (m - 1) - d
            if gap < GAP_MIN - 1e-9:
                out.append(f"Khoảng hở thép chủ cổ móng {gap * 1000:.0f} mm < {GAP_MIN * 1000:.0f} mm.")
        bg = bolt_gap(p, p["B1"])
        if bg > p["B1"] - 0.1:
            out.append(f"Cụm bu lông {bg * 1000:.0f} mm gần sát mép cổ móng {p['B1'] * 1000:.0f} mm.")
    except (KeyError, TypeError, ValueError, ZeroDivisionError):
        pass
    return out


# vùng giấy tờ 1 dành cho mặt cắt A-A + mặt bằng (mm), xem sheet1.layout
PAPER_W = 182.0          # bề rộng mặt cắt A-A (miệng hố)
PAPER_H = 182.0          # tổng chiều cao mặt cắt (H + lót) + mặt bằng (Bt + Bm + 2 lót)


def draw_bt(p, den) -> float:
    """Khoảng cách tim trụ THỂ HIỆN trên tờ 1 (m) ở tỷ lệ 1/den. Không đủ chỗ thì thu gọn (chữ kích thước vẫn
    ghi giá trị thật), tối thiểu để 2 bản đế và 2 đáy hố đào không chồng nhau."""
    k = 1000.0 / den
    B, bt = p["B"], p["BT"]
    side = SLOPE_DRAW * pit_depth(p)
    t = p["B2"] - p["B1"]                     # tim bản đế = tim đỉnh cổ + t
    w = PAPER_W / k - B - 2 * p["MPIT"] - 2 * side - t
    h = PAPER_H / k - (H(p) + TL) - B - 2 * ML - t
    lo = max(B + 0.6, B + 2 * p["MPIT"] + 0.2) - t
    return min(bt, max(lo, min(w, h)))


def auto_scale(p) -> int:
    """Tỷ lệ mặt cắt A-A và mặt bằng (tờ 1): tỷ lệ lớn nhất trong SCALES vẽ được đúng tỷ lệ (không thu gọn
    khoảng cách tim trụ). Nếu 1/250 vẫn không đủ chỗ thì dùng 1/250 và thu gọn khoảng cách tim trụ (draw_bt)."""
    for d in SCALES:
        if draw_bt(p, d) >= p["BT"] - 1e-9:
            return d
    return SCALES[-1]


def pit_joint_draw(p) -> bool:
    """Hố đào THỂ HIỆN trên hình: theo lựa chọn PIT_DRAW, mặc định theo cách tính khối lượng (pit_joint).
    Không ảnh hưởng khối lượng đào / đắp."""
    m = int(p.get("PIT_DRAW", 0) or 0)
    return pit_joint(p) if m == 0 else m == 1


def is_tkkt(p) -> bool:
    return "TKKT" in str(p.get("GD", "")).upper()


def is_detail(p) -> bool:
    """Có tờ cốt thép (tờ 2): TKBVTC và TKKT. TKKT không có bảng thống kê thép, tag không ghi số hiệu."""
    gd = str(p.get("GD", "")).upper()
    return "BVTC" in gd or "TKKT" in gd


def n_sheets(p) -> int:
    if not is_detail(p):
        return 1
    from . import sheet2
    from .geom import Geo
    return 3 if sheet2.need_sheet3(Geo(p)) else 2
