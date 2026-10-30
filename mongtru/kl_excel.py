"""Xuất bảng khối lượng ra Excel (theo app móng bản không dầm): bê tông, lót, đất, ván khuôn tính bằng công thức
tham chiếu sheet THONG SO; cốt thép lấy chiều dài và số thanh do chương trình tính theo bố trí trên bản vẽ."""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName

from . import model
from .geom import ANCHOR, CH_LEG, HOOK, KGM, Geo, tie_plan_ex

THIN = Side(style="thin", color="808080")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill("solid", fgColor="DDEBF7")
INPUT = PatternFill("solid", fgColor="FFF2CC")
B = Font(bold=True)

# (tên, ký hiệu, khoá thông số, hệ số đổi, đơn vị)
# Lưu ý: tên định danh Excel không phân biệt hoa/thường -> đặt tên khác nhau (TS_Bm, TS_b1 ...)
PARAMS = [
    ("Cạnh bản đế (Bm = Lm)", "Bm", "B", 1, "m"),
    ("Chiều sâu chôn móng", "Hm", "H1", 1, "m"),
    ("Cao cổ móng trên mặt đất", "d", "H2", 1, "m"),
    ("Khoảng cách tim trụ", "Bt", "BT", 1, "m"),
    ("Cạnh cổ móng ở đỉnh", "b1", "B1", 1, "m"),
    ("Cạnh cổ móng ở chân", "b2", "B2", 1, "m"),
    ("Cao bản đế tại gờ cổ móng", "hb", "HB", 1, "m"),
    ("Cao bản đế ở mép", "hv", "HV", 1, "m"),
    ("Gờ quanh chân cổ móng", "mpl", None, model.MPL, "m"),
    ("Dày bê tông lót", "tl", None, model.TL, "m"),
    ("Lót tràn ra mỗi bên", "ml", None, model.ML, "m"),
    ("Hệ số mở mái hố đào", "m", "SLOPE", 1, ""),
    ("Mở rộng đáy hố đào (mỗi bên)", "mp", "MPIT", 1, "m"),
    ("Số trụ móng", "n", "NPIER", 1, "trụ"),
    ("Dự phòng thép", "dp", "DP", 1, "%"),
    ("Lớp bê tông bảo vệ đáy", "cvb", "COVER_B", 1, "m"),
    ("Lớp bê tông bảo vệ còn lại", "cv", "COVER", 1, "m"),
    ("Thép bản đế lớp dưới: Ø", "db", "RB_BOT_D", 1, "mm"),
    ("Thép bản đế lớp dưới: bước", "sb", "RB_BOT_S", 1, "mm"),
    ("Thép bản đế lớp trên: Ø", "dtr", "RB_TOP_D", 1, "mm"),
    ("Thép bản đế lớp trên: bước", "str", "RB_TOP_S", 1, "mm"),
    ("Thanh chống: Ø", "dch", "RB_CH_D", 1, "mm"),
    ("Thanh chống: bước", "sch", "RB_CH_S", 1, "mm"),
    ("Thép chủ cổ móng: Ø", "dcm", "RC_D", 1, "mm"),
    ("Thép chủ cổ móng: tổng số thanh", "ncm", "RC_N", 1, "thanh"),
    ("Đai cổ móng: Ø", "dh", "RT_D", 1, "mm"),
    ("Đai cổ móng: bước", "sh", "RT_S", 1, "mm"),
    ("Đai phụ cổ móng: Ø", "dph", "RP_D", 1, "mm"),
    ("Đai phụ cổ móng: bước", "sph", "RP_S", 1, "mm"),
    ("Lưới đầu trụ: Ø", "dl", "RL_D", 1, "mm"),
    ("Lưới đầu trụ: bước", "sl", "RL_S", 1, "mm"),
    ("Lưới đầu trụ: số lớp", "nl", "RL_N", 1, "lớp"),
    ("Móc đai, đai phụ", "hk", None, HOOK, "m"),
    ("Chân thanh chống", "lch", None, CH_LEG, "m"),
    ("Neo chân thép chủ cổ móng", "lneo", None, ANCHOR, "m"),
]
NAME = {x[1]: f"TS_{x[1]}" for x in PARAMS}


def _cell(ws, r, c, v, fmt=None, bold=False, fill=None, align=None):
    x = ws.cell(row=r, column=c, value=v)
    x.border = BOX
    if fmt:
        x.number_format = fmt
    if bold:
        x.font = B
    if fill:
        x.fill = fill
    x.alignment = align or Alignment(vertical="center", wrap_text=True)
    return x


def export(p: dict, g: Geo, path: str):
    if not getattr(g, "kl", None):
        g.calc_kl()
    wb = Workbook()
    ws = wb.active
    ws.title = "THONG SO"
    for col, w in zip("ABCDE", (6, 36, 10, 12, 8)):
        ws.column_dimensions[col].width = w
    ws["A1"] = f"THÔNG SỐ MÓNG {model.type_name(p)} - CỘT {model.HT_NAMES[int(p['HT'])]} {p['COT']} ({p['GD']})"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Ô nền vàng là số liệu đầu vào - sửa ở đây, sheet KHOI LUONG tự tính lại."
    ws["A2"].font = Font(italic=True, color="808080")
    r = 4
    for i, h in enumerate(["STT", "Thông số", "Ký hiệu", "Giá trị", "Đơn vị"], 1):
        _cell(ws, r, i, h, bold=True, fill=HEAD, align=Alignment(horizontal="center", vertical="center"))
    for k, (nm, sym, key, f, unit) in enumerate(PARAMS, 1):
        r += 1
        v = p[key] * f if key else f
        _cell(ws, r, 1, k, align=Alignment(horizontal="center"))
        _cell(ws, r, 2, nm)
        _cell(ws, r, 3, sym, align=Alignment(horizontal="center"))
        _cell(ws, r, 4, v, "0.000" if unit == "m" else "0.##", fill=INPUT)
        _cell(ws, r, 5, unit, align=Alignment(horizontal="center"))
        wb.defined_names[NAME[sym]] = DefinedName(NAME[sym], attr_text=f"'THONG SO'!$D${r}")
    r += 2
    ws.cell(row=r, column=2, value="TRỌNG LƯỢNG THÉP (TCVN 1651-2018)").font = B
    r += 1
    _cell(ws, r, 2, "Đường kính (mm)", bold=True, fill=HEAD)
    _cell(ws, r, 3, "kg/m", bold=True, fill=HEAD)
    r0 = r + 1
    for d, w in sorted(KGM.items()):
        r += 1
        _cell(ws, r, 2, d)
        _cell(ws, r, 3, w, "0.000", fill=INPUT)
    wb.defined_names["TS_KGM"] = DefinedName("TS_KGM", attr_text=f"'THONG SO'!$B${r0}:$C${r}")
    # bảng tra bố trí đai phụ cổ móng theo số thanh 1 cạnh (m) và bước giữ thanh g (1 hoặc 2) - như app móng bản có dầm
    r += 2
    ws.cell(row=r, column=2, value="BỐ TRÍ ĐAI PHỤ CỔ MÓNG (chương trình lập, dùng cho sheet KHOI LUONG)").font = B
    r += 1
    heads = ["Số thanh 1 cạnh m"]
    for gg in (1, 2):
        for vt in ("cổ thẳng", "cổ vát"):
            heads += [f"g={gg}, {vt}: số đai thẳng", f"g={gg}, {vt}: chỉ số max đai bát giác",
                      f"g={gg}, {vt}: số đai xiên"]
    for i, h in enumerate(heads, 2):
        _cell(ws, r, i, h, bold=True, fill=HEAD)
    r0 = r + 1
    for m in range(2, 31):
        r += 1
        _cell(ws, r, 2, m)
        k = 3
        for gg in (1, 2):
            for vt in (False, True):
                grp, st, nd = tie_plan_ex(m, gg, vt)
                _cell(ws, r, k, len(st))
                _cell(ws, r, k + 1, max(grp) if grp else -1)
                _cell(ws, r, k + 2, nd)
                k += 3
    wb.defined_names["TS_DPC"] = DefinedName("TS_DPC", attr_text=f"'THONG SO'!$B${r0}:$N${r}")

    # ------------------------------------------------------------------ KHOI LUONG
    wk = wb.create_sheet("KHOI LUONG")
    for col, w in zip("ABCDEFGHIJ", (6, 44, 10, 70, 12, 9, 12, 12, 34, 40)):
        wk.column_dimensions[col].width = w
    wk["A1"] = f"BẢNG KHỐI LƯỢNG CHO 01 VỊ TRÍ MÓNG {model.type_name(p)} - CỘT {p['COT']}"
    wk["A1"].font = Font(bold=True, size=12)
    r = 3
    for i, h in enumerate(["STT", "Hạng mục", "Ký hiệu", "Công thức / diễn giải", "Khối lượng", "Đơn vị"], 1):
        _cell(wk, r, i, h, bold=True, fill=HEAD, align=Alignment(horizontal="center", vertical="center"))
    rows = {}

    def line(stt, name, sym, formula, desc, unit, bold=False):
        nonlocal r
        r += 1
        _cell(wk, r, 1, stt, align=Alignment(horizontal="center"))
        _cell(wk, r, 2, name, bold=bold)
        _cell(wk, r, 3, sym, align=Alignment(horizontal="center"))
        _cell(wk, r, 4, desc)
        _cell(wk, r, 5, formula, "0.00", bold=bold)
        _cell(wk, r, 6, unit, align=Alignment(horizontal="center"))
        if sym:
            rows[sym] = f"$E${r}"
        return f"$E${r}"

    def sec(t):
        nonlocal r
        r += 1
        c = _cell(wk, r, 1, t, bold=True, fill=HEAD)
        wk.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        c.alignment = Alignment(vertical="center")

    sec("A. KÍCH THƯỚC TRUNG GIAN")
    a = line("", "Cạnh gờ quanh chân cổ móng", "a", "=TS_b2+2*TS_mpl", "b2 + 2·mpl", "m")
    hc = line("", "Chiều cao cổ móng (từ mặt gờ)", "hc", "=TS_Hm+TS_d-TS_hb", "Hm + d − hb", "m")
    hd = line("", "Chiều sâu hố đào", "Hd", "=TS_Hm+TS_tl", "Hm + dày lót", "m")
    dd = line("", "Khoảng cách tim bản đế", "D", "=TS_Bt+TS_b2-TS_b1",
              "Bt + (b2 − b1): cổ vát ra phía ngoài, mặt trong thẳng đứng", "m")
    kh = line("", "Khoảng hở giữa miệng 2 hố đào riêng", "kh", f"={dd}-2*(0.5*TS_Bm+TS_mp+{hd}*TS_m)",
              "D − 2·(Bm/2 + mp + Hd·m); <= 0 thì đào chung 1 hố cho 4 trụ", "m")
    a1 = line("", "Cạnh đáy hố đào", "a1", f"=IF({kh}>0,TS_Bm+2*TS_mp,{dd}+TS_Bm+2*TS_mp)",
              "Hố riêng: Bm + 2·mp; hố chung: D + Bm + 2·mp", "m")
    a2 = line("", "Cạnh miệng hố đào", "a2", f"={a1}+2*TS_m*{hd}", "a1 + 2·m·Hd", "m")
    nh = line("", "Số hố đào", "nh", f"=IF({kh}>0,TS_n,TS_n/4)", "Hố riêng: n; hố chung: n/4", "hố")
    bt_ = line("", "Cạnh cổ móng tại mặt đất", "bd", "=TS_b2+(TS_b1-TS_b2)*(TS_Hm-TS_hb)/(TS_Hm+TS_d-TS_hb)",
               "b2 + (b1 − b2)·(Hm − hb)/hc", "m")

    sec("B. BÊ TÔNG")
    vb = line(1, f"Bê tông móng {p['BTM']}", "Vbt",
              f"=TS_n*(TS_Bm^2*TS_hv+(TS_Bm^2+{a}^2+TS_Bm*{a})*(TS_hb-TS_hv)/3+(TS_b1^2+TS_b2^2+TS_b1*TS_b2)*{hc}/3)",
              "n·[Bm²·hv + (Bm² + a² + Bm·a)·(hb − hv)/3 + (b1² + b2² + b1·b2)·hc/3]", "m3", True)
    vl = line(2, f"Bê tông lót {p['BTL']}", "Vlot", "=TS_n*(TS_Bm+2*TS_ml)^2*TS_tl", "n·(Bm + 2·ml)²·tl", "m3", True)

    sec("C. ĐẤT, VÁN KHUÔN")
    vd = line(3, "Đào đất hố móng", "Vdao", f"={nh}*{hd}/3*({a1}^2+{a2}^2+{a1}*{a2})",
              "nh · Hd/3·(a1² + a2² + a1·a2)", "m3", True)
    va = line("", "Thể tích cổ móng trên mặt đất", "Vtr", f"=TS_n*TS_d/3*({bt_}^2+TS_b1^2+{bt_}*TS_b1)",
              "n·d/3·(bd² + b1² + bd·b1)", "m3")
    vdp = line(4, "Đắp đất hoàn trả hố móng", "Vdap", f"={vd}-{vl}-({vb}-{va})", "Vdao − Vlot − (Vbt − Vtr)", "m3", True)
    line(5, "Đất thừa", "Vthua", f"={vd}-{vdp}", "Vdao − Vdap", "m3", True)
    line(6, "Ván khuôn móng", "VK",
         f"=TS_n*((TS_b1+TS_b2)*({hc}+SQRT({hc}^2+(TS_b2-TS_b1)^2))+4*TS_Bm*TS_hv)",
         "n·[2 mặt cổ đứng + 2 mặt cổ vát: (b1 + b2)·(hc + √(hc² + (b2 − b1)²)) + 4·Bm·hv] "
         "(mặt vát bản đế không ghép ván khuôn)", "m2", True)

    sec("D. CỐT THÉP - ĐẠI LƯỢNG TRUNG GIAN (1 trụ)")
    E = 0.000001
    sp = line("", "Chiều dài rải thép bản đế", "sp", "=TS_Bm-2*TS_cv", "Bm − 2cv", "m")
    pm = line("", "Nửa gờ phẳng quanh chân cổ móng", "pm", "=TS_b2/2+TS_mpl", "b2/2 + mpl", "m")
    hl = line("", "Nửa chiều dài rải", "hl", f"={sp}/2", "sp/2", "m")

    def zone(sym, s_, nm):
        nf = line("", f"{nm}: số thanh vùng gờ (1 phương)", f"nf{sym}",
                  f"=IF({pm}>={hl}-{E},MAX(1,ROUNDUP({sp}/({s_}/1000)-{E},0)),MAX(1,ROUNDUP(2*{pm}/({s_}/1000)-{E},0)))",
                  "ROUNDUP(2·pm/@) (rải vùng gờ, thanh cách mép vùng 1/2 bước)", "thanh")
        ns = line("", f"{nm}: số thanh vùng vát mỗi bên", f"ns{sym}",
                  f"=IF({pm}>={hl}-{E},0,MAX(1,ROUNDUP(({hl}-{pm})/({s_}/1000)-{E},0)))", "ROUNDUP((hl − pm)/@)", "thanh")
        return nf, ns
    nf1, ns1 = zone("1", "TS_sb", "Lớp dưới")
    nf2, ns2 = zone("2", "TS_str", "Lớp trên")
    ee = line("", "Mút thanh lớp trên (tính từ tim)", "e", "=TS_Bm/2-TS_cv", "Bm/2 − cv", "m")
    hk2 = line("", "Chân móc 2 đầu thanh lớp trên", "hk2",
               f"=MAX(0,ROUNDDOWN((TS_hv-TS_cv-TS_cvb-(TS_db+TS_dtr)/1000)/0.005+{E},0)*0.005)",
               "hv − cv − cvb − (Ø dưới + Ø trên), làm tròn xuống tới 5 mm", "m")
    kv = line("", "Hệ số dài đoạn xiên", "kv", f"=SQRT(1+((TS_hb-TS_hv)/(TS_Bm/2-{pm}))^2)",
              "√(1 + ((hb − hv)/(Bm/2 − pm))²)", "")
    d0 = line("", "Đoạn xiên thanh trong vùng gờ", "d0", f"=({ee}-{pm})*{kv}", "(e − pm)·kv", "m")
    li = line("", "Chiều dài thanh lớp trên trong vùng gờ", "Li",
              f"=IF({d0}<0.06,2*{ee},2*{pm}+2*{d0})+2*{hk2}", "2·pm + 2·d0 + 2·móc (d0 < 60 thì thẳng 2e)", "m")
    ss2 = line("", "Bước thanh lớp trên vùng vát", "ss2", f"=IF({ns2}>0,({hl}-{pm})/{ns2},0)", "(hl − pm)/ns2", "m")
    J = f"(ROW(OFFSET($A$1,0,0,{ns2},1))-0.5)"
    aj = f"({pm}+{J}*{ss2})"
    dj = f"(({ee}-{aj})*{kv})"
    so = line("", "Tổng chiều dài thanh lớp trên vùng vát (1 bên, 1 phương)", "So",
              f"=IF({ns2}>0,SUMPRODUCT(({dj}<0.06)*2*{ee}+({dj}>=0.06)*(2*{aj}+2*{dj})+2*{hk2}),0)",
              "Thanh j cách tim aj = pm + (j − 0.5)·ss2: dài 2·aj + 2·(e − aj)·kv + 2·móc", "m")
    lch = line("", "Chiều dài rải thanh chống", "lc", "=MAX(TS_sch/1000,{0}-2*TS_sch/1000)".format(sp),
               "MAX(@, sp − 2@)", "m")
    nch = line("", "Số thanh chống mỗi phương", "nch", f"=1+MAX(1,ROUNDUP({lch}/(TS_sch/1000)-{E},0))",
               "1 + ROUNDUP(lc/@)", "thanh")
    X = f"ABS({lch}*((ROW(OFFSET($A$1,0,0,{nch},1))-1)/({nch}-1)-0.5))"
    Y = f"ABS({lch}*((COLUMN(OFFSET($A$1,0,0,1,{nch}))-1)/({nch}-1)-0.5))"
    xnote = ""
    if g.chair_rot():
        # thanh chống xoay: thân đặt sát thép lưới phương Y (Geo.chair_grid) - |x| từng thanh theo bố trí
        cg = {(i, j): abs(x) for i, j, x, _ in g.chair_grid()}
        ni_ = 1 + max(i for i, _ in cg)
        X = "{" + ";".join(",".join(f"{cg[i, j]:.6f}" for j in range(ni_)) for i in range(ni_)) + "}"
        xnote = "; thanh chống xoay: |x| từng thanh = vị trí sát cạnh thép lưới phương Y gần lưới @ nhất (phía tim)"
    R = f"(({X}>={Y})*{X}+({X}<{Y})*{Y})"
    sch = line("", "Tổng chiều cao thân thanh chống (1 trụ)", "Sch",
               f"=SUMPRODUCT(TS_hb-(TS_hb-TS_hv)*({R}>{pm})*({R}-{pm})/(TS_Bm/2-{pm}))"
               f"-{nch}^2*(TS_cv+TS_cvb+(TS_db+TS_dtr)/1000)",
               "Σ thân = cao bản đế tại vị trí thanh − 2 lớp bảo vệ (cv + cvb) − 2Ø (Ø lớp dưới + Ø lớp trên); "
               "cao bản đế theo max(|x|,|y|) của thanh" + xnote, "m")
    hh_ = line("", "Chiều cao toàn trụ", "H", "=TS_Hm+TS_d", "Hm + d", "m")
    z0 = line("", "Cao độ đai dưới lớp lưới đầu trụ thấp nhất", "z0",
              f"={hh_}-TS_cv-(TS_nl-1)*TS_sl/1000-2*TS_dl/1000", "H − cv − (số lớp − 1)·@lưới − 2Ø lưới", "m")
    n5 = line("", "Số đai cổ móng (1 trụ)", "n5",
              f"=TS_nl+IF({z0}-TS_hb<={E},0,MAX(1,ROUNDUP(({z0}-TS_hb)/(TS_sh/1000)-{E},0)))",
              "Số lớp lưới (đai đỡ lưới) + ROUNDUP((z0 − hb)/@)", "đai")
    nt = line("", "Số mức đai phụ cổ móng (1 trụ)", "nt",
              f"=1+MAX(1,ROUNDUP(({hh_}-TS_cv-2*TS_dl/1000-TS_hb)/(TS_sph/1000)-{E},0))",
              "1 + ROUNDUP((H − cv − 2Ø lưới − hb)/@ đai phụ)", "mức")
    # cổ móng vát: cạnh cổ b(z) = b2 − kz·(z − hb); đai / đai phụ có chiều dài tuyến tính theo b -> dùng cao độ TB
    kz = line("", "Độ vát cổ móng (cạnh giảm theo chiều cao)", "kz", f"=(TS_b2-TS_b1)/({hh_}-TS_hb)",
              "(b2 − b1)/(H − hb); 0 khi cổ thẳng", "m/m")
    nl5 = f"({n5}-TS_nl)"
    z5 = line("", "Cao độ trung bình các đai cổ móng", "z5",
              f"=(TS_nl*({hh_}-TS_cv-2*TS_dl/1000)-TS_sl/1000*TS_nl*(TS_nl-1)/2"
              f"+IF({nl5}>0,{nl5}*{z0}-({z0}-TS_hb)*({nl5}+1)/2,0))/{n5}",
              "Đai đỡ lưới: H − cv − 2Ø lưới − i·@lưới; còn lại rải đều từ z0 xuống hb", "m")
    b5 = line("", "Cạnh cổ móng trung bình tại các đai", "b5", f"=TS_b2-{kz}*({z5}-TS_hb)", "b2 − kz·(z5 − hb)", "m")
    bt7 = line("", "Cạnh cổ móng trung bình tại các mức đai phụ", "b7",
               f"=TS_b2-{kz}*(({hh_}-TS_cv-2*TS_dl/1000+TS_hb)/2-TS_hb)",
               "Mức đai phụ rải đều từ H − cv − 2Ø lưới xuống hb: cao độ TB = (đỉnh + hb)/2", "m")
    vc = line("", "Chiều cao đoạn thẳng đứng thép chủ cổ móng", "Vc",
              f"={hh_}-TS_cv-TS_cvb-2*TS_db/1000-TS_dcm/2000",
              "H − cv − cvb − 2Ø lớp dưới − Ø chủ/2 (đoạn neo tựa trên 2 phương thép lớp dưới)", "m")
    mcm = line("", "Số thanh thép chủ cổ móng 1 cạnh", "mcm", "=1+INT(TS_ncm/4)", "1 + n/4", "thanh")
    scm = line("", "Khoảng cách thanh chủ cổ móng", "scm", f"=1000*(TS_b1-2*TS_cv)/MAX(1,{mcm}-1)",
               "(b − 2cv)/(m − 1)", "mm")
    gst = line("", "Bước thanh được đai phụ giữ", "gst", f"=MAX(1,MIN(2,INT(400/{scm})))",
               "Khoảng cách thanh được giữ <= 400 mm: 1 (mọi thanh) hoặc 2 (cách 1 thanh)", "")
    none_ = f"OR({mcm}<=2,AND(1000*(TS_b1-2*TS_cv)<=400,{mcm}<=4))"
    col = f"(2+6*({gst}-1)+3*(TS_b2>TS_b1+{E}))"
    nst = line("", "Số đai phụ thẳng mỗi mức (mỗi phương)", "nst",
               f"=IF({none_},0,VLOOKUP({mcm},TS_DPC,{col},FALSE))",
               "Tra bảng BỐ TRÍ ĐAI PHỤ CỔ MÓNG (sheet THONG SO) theo m, g và cổ thẳng / vát", "thanh")
    gmax = line("", "Chỉ số thanh xa nhất của đai bát giác (−1: không có)", "gmax",
                f"=IF({none_},-1,VLOOKUP({mcm},TS_DPC,{col}+1,FALSE))", "Tra bảng như trên", "")
    ndg = line("", "Số đai xiên mỗi mức (mỗi phương)", "ndg",
               f"=IF({none_},0,VLOOKUP({mcm},TS_DPC,{col}+2,FALSE))",
               "Cổ vát, số thanh 1 cạnh chẵn: đai thẳng đặt xiên nối 2 thanh giữa của 2 cạnh đối diện", "")
    zt_ = f"({hh_}-TS_cv-2*TS_dl/1000)"
    K_ = f"(ROW(OFFSET($A$1,0,0,{nt},1))-1)"
    bk = f"(TS_b2-{kz}*({zt_}-{K_}*({zt_}-TS_hb)/({nt}-1)-TS_hb))"
    sdg = line("", "Tổng thân đai xiên các mức (1 phương)", "Sdg",
               f"=SUMPRODUCT(SQRT(({bk}-2*TS_cv)^2+(({bk}-2*(TS_cv+TS_dh/1000+TS_dcm/2000))/({mcm}-1))^2))",
               "Mức k: b = cạnh cổ tại mức đai; thân = √((b − 2cv)² + s²), s = (b − 2(cv + Ø đai + Ø chủ/2))/(m − 1)",
               "m")
    rr = "MAX(0.35,0.5*TS_dcm/1000*40)"
    hb_ = f"(40*({bt7}/2-TS_cv)-0.2-{rr})"
    rt = f"({rr}+0.12)"
    oa = line("", "Đai bát giác: nửa cạnh thẳng", "ao",
              f"=IF({gmax}<0,0,(({gmax}-({mcm}-1)/2)*2*{hb_}/({mcm}-1)+0.41421356*{rt})/40)",
              "Theo hình mặt cắt 1-1 (bo theo thép chủ), cạnh cổ b7", "m")
    of = line("", "Đai bát giác: nửa bề rộng", "fo", f"=IF({gmax}<0,0,({hb_}+{rt})/40)", "Theo hình mặt cắt 1-1", "m")
    nn8 = line("", "Số thanh lưới đầu trụ mỗi phương, mỗi lớp", "n8",
               f"=1+ROUNDUP((TS_b1-2*TS_cv)/(TS_sl/1000)-{E},0)", "1 + ROUNDUP((b − 2cv)/@)", "thanh")

    I_ = f"(ROW(OFFSET($A$1,0,0,{mcm},1))-1)"
    J_ = f"(COLUMN(OFFSET($A$1,0,0,1,{mcm}))-1)"
    msk = f"((({I_}=0)+({I_}={mcm}-1)+({J_}=0)+({J_}={mcm}-1))>0)"
    qv = f"({kz}*{vc}/({mcm}-1))"
    s4 = line("", "Tổng chiều dài đoạn thân thép chủ (1 trụ)", "S4",
              f"=SUMPRODUCT({msk}*SQRT({vc}^2+({I_}*{qv})^2+({J_}*{qv})^2))",
              "Thanh (i, j) trên chu vi (i, j tính từ 2 mặt đứng phía trong): √(Vc² + (i·q)² + (j·q)²), "
              "q = kz·Vc/(m − 1)", "m")
    # công thức từng vị trí: (Ø, chiều dài 1 thanh (m), số lượng cả móng, diễn giải chiều dài, diễn giải số lượng)
    n2 = f"({nf2}+2*{ns2})"
    F = {
        1: ("TS_db", f"{sp}", f"2*({nf1}+2*{ns1})*TS_n", "Bm − 2cv", "2 phương × (nf1 + 2·ns1) × số trụ"),
        2: ("TS_dtr", f"({nf2}*{li}+2*{so})/{n2}", f"2*{n2}*TS_n",
            "Trung bình: (nf2·Li + 2·So)/(nf2 + 2·ns2)", "2 phương × (nf2 + 2·ns2) × số trụ"),
        3: ("TS_dch", f"{sch}/{nch}^2+2*TS_lch", f"{nch}^2*TS_n", "Trung bình: Sch/nch² + 2 chân 100",
            "nch² × số trụ"),
        4: ("TS_dcm", f"{s4}/(4*({mcm}-1))+TS_lneo", "TS_ncm*TS_n",
            "Trung bình: S4/số thanh + neo 500 (cổ thẳng: Vc + 500)", "số thanh × số trụ"),
        5: ("TS_dh", f"4*({b5}-2*TS_cv)+2*TS_hk", f"{n5}*TS_n", "Trung bình: 4(b5 − 2cv) + 2 móc", "n5 × số trụ"),
        6: ("TS_dph", f"IF({gmax}<0,0,8*{oa}+4*SQRT(2)*({of}-{oa})+2*TS_hk)", f"IF({gmax}<0,0,{nt}*TS_n)",
            "8·ao + 4√2(fo − ao) + 2 móc", "nt × số trụ"),
        7: ("TS_dph", f"{bt7}-2*TS_cv+2*TS_hk", f"{nt}*2*{nst}*TS_n", "Trung bình: b7 − 2cv + 2 móc",
            "nt × 2 phương × nst × số trụ"),
        9: ("TS_dph", f"{sdg}/{nt}+2*TS_hk", f"{nt}*2*{ndg}*TS_n", "Trung bình: Sdg/nt + 2 móc",
            "nt × 2 phương × ndg × số trụ"),
        8: ("TS_dl", "TS_b1-2*TS_cv", f"TS_nl*2*{nn8}*TS_n", "b − 2cv", "số lớp × 2 phương × n8 × số trụ"),
    }
    from .geom import NAMES, POS_ORDER
    sec("E. CỐT THÉP (chiều dài, số lượng tính bằng công thức - xem cột I, J)")
    r += 1
    for i, h in enumerate(["Số hiệu", "Hình dạng / ghi chú", "Ø (mm)", "Chiều dài 1 thanh (mm)", "Số lượng",
                           "kg/m", "Tổng dài (m)", "KL (kg)", "Cách tính chiều dài", "Cách tính số lượng"], 1):
        _cell(wk, r, i, h, bold=True, fill=HEAD, align=Alignment(horizontal="center", vertical="center", wrap_text=True))
    rs0 = r + 1
    for pos in POS_ORDER:
        r += 1
        fd, fl, fq, tl, tq = F[pos]
        # số hiệu chạy theo số liệu: chỉ đánh số các thanh có số lượng > 0 (giống bảng thống kê trên bản vẽ)
        _cell(wk, r, 1, f'=IF(E{r}>0,COUNTIF(E${rs0}:E{r},">0"),"—")', align=Alignment(horizontal="center"))
        _cell(wk, r, 2, NAMES[pos] + (" - chiều dài trung bình" if pos in (2, 3) else ""))
        _cell(wk, r, 3, f"={fd}", "0")
        _cell(wk, r, 4, f"=ROUND(1000*({fl}),0)", "0")
        _cell(wk, r, 5, f"={fq}", "0")
        _cell(wk, r, 6, f"=IFERROR(VLOOKUP(C{r},TS_KGM,2,FALSE),C{r}^2/162)", "0.000")
        _cell(wk, r, 7, f"=D{r}/1000*E{r}", "0.00")
        _cell(wk, r, 8, f"=G{r}*F{r}", "0.00")
        _cell(wk, r, 9, tl)
        _cell(wk, r, 10, tq)
    rs1 = r
    r += 1
    _cell(wk, r, 2, "TỔNG CỘNG THÉP", bold=True)
    _cell(wk, r, 8, f"=SUM(H{rs0}:H{rs1})", "0.00", bold=True)
    r += 1
    _cell(wk, r, 2, "TỔNG CỘNG THÉP (có dự phòng)", bold=True)
    _cell(wk, r, 4, "Tổng cộng × (1 + dp/100)")
    _cell(wk, r, 8, f"=H{r - 1}*(1+TS_dp/100)", "0.00", bold=True)
    rtot = r
    r += 2
    _cell(wk, r, 2, "Phân loại theo đường kính (có dự phòng)", bold=True, fill=HEAD)
    _cell(wk, r, 3, "Ø (mm)", bold=True, fill=HEAD)
    _cell(wk, r, 7, "Tổng dài (m)", bold=True, fill=HEAD)
    _cell(wk, r, 8, "KL (kg)", bold=True, fill=HEAD)
    rdia = {}
    for d in sorted(KGM):                   # đủ các Ø trong bảng tra (đổi Ø ở THONG SO vẫn có dòng)
        r += 1
        rdia[int(d + 0.5)] = r
        _cell(wk, r, 2, f"Thép Ø{int(d + 0.5)}")
        _cell(wk, r, 3, int(d + 0.5))
        _cell(wk, r, 7, f"=SUMIF($C${rs0}:$C${rs1},C{r},$G${rs0}:$G${rs1})*(1+TS_dp/100)", "0.00")
        _cell(wk, r, 8, f"=SUMIF($C${rs0}:$C${rs1},C{r},$H${rs0}:$H${rs1})*(1+TS_dp/100)", "0.00")
    r += 2
    _cell(wk, r, 2, "Bu lông neo", bold=True)
    _cell(wk, r, 4, f"{model.bolt_text(p)}  ({p['BLN']})")
    wk.freeze_panes = "A4"
    used = sorted({int(d + 0.5) for d, _, _ in g.klc})
    _summary(wb, p, rows, {d: rdia[d] for d in used if d in rdia}, rtot)
    wb.active = 0
    wb.save(path)


def _summary(wb, p, rows, rdia, rtot):
    """Sheet TONG KE: bảng tổng kê khối lượng 1 vị trí móng, mỗi số liệu 1 ô."""
    ws = wb.create_sheet("TONG KE", 0)
    K = "'KHOI LUONG'!"
    cols = [(f"{p['BTM']} (m³)", f"={K}{rows['Vbt']}", "0.00"),
            (f"{p['BTL']} (m³)", f"={K}{rows['Vlot']}", "0.00")]
    for d in sorted(rdia):
        cols.append((f"Ø{d} (kg)", f"={K}$H${rdia[d]}", "0.00"))
    cols += [("TỔNG K.LƯỢNG THÉP (kg)", f"={K}$H${rtot}", "0.00"),
             ("BULONG NEO", model.bolt_text(p), None),
             ("Đào đất (m³)", f"={K}{rows['Vdao']}", "0.00"),
             ("Đắp đất (m³)", f"={K}{rows['Vdap']}", "0.00"),
             ("Ván khuôn (m²)", f"={K}{rows['VK']}", "0.00"),
             ("Rộng đền bù (m)", "=TS_Bt+TS_b2-TS_b1+TS_Bm+2*0.5", "0.00"),
             ("Dài đền bù (m)", "=TS_Bt+TS_b2-TS_b1+TS_Bm+2*0.5", "0.00"),
             ("Diện đền bù (m²)", None, "0.00"),
             ("Chiều sâu hố đào (m)", f"={K}{rows['Hd']}", "0.00")]
    n = len(cols)
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 64
    ws.row_dimensions[3].height = 24
    c = ws.cell(row=1, column=1, value="BẢNG TỔNG KÊ KHỐI LƯỢNG CHO 01 VỊ TRÍ MÓNG")
    c.font = Font(size=14)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    for j in range(1, n + 1):
        ws.cell(row=1, column=j).border = BOX
    from openpyxl.utils import get_column_letter as L
    for j, (h, f, fmt) in enumerate(cols, 1):
        ws.column_dimensions[L(j)].width = 11 if len(h) < 14 else 12.5
        _cell(ws, 2, j, h, align=Alignment(horizontal="center", vertical="center", wrap_text=True))
        if h.startswith("Diện đền bù"):
            f = f"={L(j - 2)}3*{L(j - 1)}3"
        _cell(ws, 3, j, f, fmt, align=Alignment(horizontal="center", vertical="center"))
    if p.get("DP", 0) > 0:                  # dự phòng = 0: không ghi chú
        ws.cell(row=5, column=1, value=f"Khối lượng thép đã gồm {p.get('DP', 0):g}% dự phòng.").font = \
            Font(italic=True, color="808080")
    # kiểm tra: đổi Ø ở THONG SO sang đường kính không có cột ở đây thì tổng các cột Ø lệch tổng thép
    jd = [j for j, (h, _, _) in enumerate(cols, 1) if h.startswith("Ø")]
    jt = next(j for j, (h, _, _) in enumerate(cols, 1) if h.startswith("TỔNG K.LƯỢNG THÉP"))
    sm = "+".join(f"{L(j)}3" for j in jd) or "0"
    c = ws.cell(row=6, column=1, value=f'=IF(ABS(({sm})-{L(jt)}3)>0.01,"Có đường kính mới - xuất lại file","")')
    c.font = Font(bold=True, color="C00000")
