"""Excel khối lượng móng 2 loại (tách riêng từng loại).

Mỗi loại: sheet "THONG SO i" + "KHOI LUONG i" - đúng như file khối lượng của bản 4 trụ giống nhau (xuất bằng
mongtru.kl_excel rồi đổi tên sheet, tên định danh TS_x -> Ti_x), số trụ = số trụ của loại đó; phần đào / đắp tính
chung cả vị trí móng ở sheet "DAT". Sheet "TONG KE": tổng kê cả vị trí móng.
"""
from __future__ import annotations

import copy
import os
import re
import tempfile

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName

from .. import kl_excel as K
from .. import model
from . import model as lm

EARTH = ("kh", "a1", "a2", "nh", "Vdao", "Vdap", "Vthua")       # dòng đất của từng loại -> tính chung ở DAT


def _copy_sheet(src, dst, ren):
    for row in src.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and v.startswith("="):
                v = ren(v)
            d = dst.cell(row=c.row, column=c.column, value=v)
            if c.has_style:
                d.font = copy.copy(c.font)
                d.border = copy.copy(c.border)
                d.fill = copy.copy(c.fill)
                d.number_format = c.number_format
                d.alignment = copy.copy(c.alignment)
    for rg in src.merged_cells.ranges:
        dst.merge_cells(str(rg))
    for k, dim in src.column_dimensions.items():
        dst.column_dimensions[k].width = dim.width
    for k, dim in src.row_dimensions.items():
        if dim.height:
            dst.row_dimensions[k].height = dim.height
    dst.freeze_panes = src.freeze_panes


def _rows(ws):
    """Ký hiệu cột C -> dòng (sheet KHOI LUONG)."""
    return {ws.cell(r, 3).value: r for r in range(1, ws.max_row + 1) if isinstance(ws.cell(r, 3).value, str)}


def export(mp: dict, gs: list, path: str):
    wb = Workbook()
    wb.remove(wb.active)
    info = []
    tmp = tempfile.mkdtemp()
    for i, g in enumerate(gs):
        t = mp["TYPES"][i]
        p = lm.type_params(mp, i)
        f = os.path.join(tmp, f"t{i}.xlsx")
        K.export(p, g, f)
        src = openpyxl.load_workbook(f)
        n = i + 1
        sts, skl = f"THONG SO {n}", f"KHOI LUONG {n}"

        def ren(s, n=n, sts=sts, skl=skl):
            s = re.sub(r"\bTS_(\w+)", rf"T{n}_\1", s)
            return s.replace("'THONG SO'!", f"'{sts}'!").replace("'KHOI LUONG'!", f"'{skl}'!")
        a = wb.create_sheet(sts)
        _copy_sheet(src["THONG SO"], a, ren)
        b = wb.create_sheet(skl)
        _copy_sheet(src["KHOI LUONG"], b, ren)
        for nm, dn in list(src.defined_names.items()):
            new = re.sub(r"^TS_", f"T{n}_", nm)
            wb.defined_names[new] = DefinedName(new, attr_text=ren(dn.attr_text))
        pos = ", ".join(str(q) for q in t["POS"])
        a["A1"] = (f"THÔNG SỐ MÓNG LOẠI {n}: {len(t['POS'])}{t['NAME']} (VỊ TRÍ TRỤ {pos}) - CỘT "
                   f"{model.HT_NAMES[int(mp['HT'])]} {mp['COT']} ({mp['GD']})")
        # cao độ thiết kế (mặt đất tại trụ) - dùng cho đào / đắp ở sheet DAT
        r = a.max_row + 2
        a.cell(row=r, column=2, value="CAO ĐỘ (dùng cho sheet DAT)").font = Font(bold=True)
        r += 1
        K._cell(a, r, 2, "Cao độ thiết kế (mặt đất tại trụ)")
        K._cell(a, r, 3, "cd", align=Alignment(horizontal="center"))
        K._cell(a, r, 4, t["CD"], "0.000", fill=K.INPUT)
        K._cell(a, r, 5, "m", align=Alignment(horizontal="center"))
        wb.defined_names[f"T{n}_cd"] = DefinedName(f"T{n}_cd", attr_text=f"'{sts}'!$D${r}")
        r += 1
        K._cell(a, r, 2, "Vị trí trụ")
        K._cell(a, r, 4, pos)
        b["A1"] = (f"BẢNG KHỐI LƯỢNG MÓNG LOẠI {n}: {len(t['POS'])}{t['NAME']} (VỊ TRÍ TRỤ {pos}) - "
                   f"CỘT {mp['COT']}")
        rows = _rows(b)
        for sym in EARTH:
            if sym in rows:
                rr = rows[sym]
                b.cell(rr, 4).value = "Đất đào / đắp tính chung cả vị trí móng: xem sheet DAT"
                b.cell(rr, 5).value = None
        # thép từng Ø (có dự phòng) và tổng thép
        rd, rtot = {}, None
        for rr in range(1, b.max_row + 1):
            v = b.cell(rr, 2).value
            if isinstance(v, str):
                m = re.fullmatch(r"Thép Ø(\d+)", v)
                if m:
                    rd[int(m.group(1))] = rr
                elif v == "TỔNG CỘNG THÉP (có dự phòng)":
                    rtot = rr
        info.append(dict(n=n, skl=skl, rows=rows, rd=rd, rtot=rtot, g=g))
    dat = _earth(wb, mp, info)
    _summary(wb, mp, info, dat)
    wb.active = 0
    wb.save(path)


def _earth(wb, mp, info):
    """Sheet DAT: đào / đắp cả vị trí móng."""
    ws = wb.create_sheet("DAT")
    for col, w in zip("ABCDEF", (6, 44, 10, 74, 12, 9)):
        ws.column_dimensions[col].width = w
    ws["A1"] = f"ĐÀO / ĐẮP ĐẤT CHO 01 VỊ TRÍ MÓNG {lm.type_name(mp)} - CỘT {mp['COT']}"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("Mỗi trụ: hố từ đáy lót (cao độ thiết kế − Hm − tl) lên mặt đất tại trụ (cao độ thiết kế), mái theo "
                "hệ số m. Miệng hố 2 trụ cạnh nhau chạm nhau thì đào chung 1 hố, đáy bậc theo độ sâu từng trụ.")
    ws["A2"].font = Font(italic=True, color="808080")
    r = 3
    for i, h in enumerate(["STT", "Hạng mục", "Ký hiệu", "Công thức / diễn giải", "Khối lượng", "Đơn vị"], 1):
        K._cell(ws, r, i, h, bold=True, fill=K.HEAD, align=Alignment(horizontal="center", vertical="center"))
    out = {}

    def line(stt, name, sym, f, desc, unit, bold=False):
        nonlocal r
        r += 1
        K._cell(ws, r, 1, stt, align=Alignment(horizontal="center"))
        K._cell(ws, r, 2, name, bold=bold)
        K._cell(ws, r, 3, sym, align=Alignment(horizontal="center"))
        K._cell(ws, r, 4, desc)
        K._cell(ws, r, 5, f, "0.000", bold=bold)
        K._cell(ws, r, 6, unit, align=Alignment(horizontal="center"))
        out[sym] = f"$E${r}"
        return f"$E${r}"
    hd, a1, a2, pc = {}, {}, {}, {}
    for x in info:
        n = x["n"]
        T = f"T{n}_"
        t = mp["TYPES"][n - 1]
        hd[n] = line("", f"Loại {n} ({t['NAME']}): chiều sâu hố", f"Hd{n}", f"={T}Hm+{T}tl", "Hm + tl", "m")
        a1[n] = line("", f"Loại {n}: cạnh đáy hố riêng", f"a1{n}", f"={T}Bm+2*{T}mp", "Bm + 2·mp", "m")
        a2[n] = line("", f"Loại {n}: cạnh miệng hố riêng", f"a2{n}", f"={a1[n]}+2*{T}m*{hd[n]}", "a1 + 2·m·Hd", "m")
        pc[n] = line("", f"Loại {n}: tâm bản đế cách tim móng (mỗi phương)", f"P{n}", f"=({T}Bt+{T}b2-{T}b1)/2",
                     "(Bt + b2 − b1)/2", "m")
    ta = lm.type_at(mp)
    gaps = []
    for qa, qb in ((1, 2), (3, 4), (1, 3), (2, 4)):
        na, nb = ta[qa] + 1, ta[qb] + 1
        gaps.append(line("", f"Khe hở miệng hố trụ {qa} - trụ {qb}", f"k{qa}{qb}",
                         f"={pc[na]}+{pc[nb]}-({a2[na]}+{a2[nb]})/2",
                         f"P{na} + P{nb} − (a2{na} + a2{nb})/2; <= 0: các hố chạm nhau", "m"))
    joint = lm.pit_joint(mp)
    if joint:
        vd = line(1, "Đào đất hố móng (đào chung 1 hố, đáy bậc)", "Vdao", round(lm.vol_dao(mp), 3),
                  "Chương trình tính theo lát cắt: diện tích hố chung tại mỗi cao độ = hình chữ nhật bao các hố "
                  "riêng có ở cao độ đó; sửa thông số thì xuất lại file", "m3", True)
        chk = f'=IF(MIN({",".join(gaps)})>0,"Các hố không còn chạm nhau - xuất lại file","")'
    else:
        terms = "+".join(f"T{x['n']}_n*{hd[x['n']]}/3*({a1[x['n']]}^2+{a2[x['n']]}^2+{a1[x['n']]}*{a2[x['n']]})"
                         for x in info)
        vd = line(1, "Đào đất hố móng (đào riêng từng trụ)", "Vdao", f"={terms}",
                  "Σ số trụ · Hd/3·(a1² + a2² + a1·a2) của từng loại", "m3", True)
        chk = f'=IF(MIN({",".join(gaps)})<=0,"Các hố chạm nhau (đào chung) - xuất lại file","")'
    kl = [f"'{x['skl']}'!" for x in info]
    vlot = "+".join(f"{k}$E${x['rows']['Vlot']}" for k, x in zip(kl, info))
    vbt = "+".join(f"({k}$E${x['rows']['Vbt']}-{k}$E${x['rows']['Vtr']})" for k, x in zip(kl, info))
    vdp = line(2, "Đắp đất hoàn trả hố móng", "Vdap", f"={vd}-({vlot})-({vbt})",
               "Vdao − Σ Vlot − Σ (Vbt − Vtr) của 2 loại (sheet KHOI LUONG 1, 2)", "m3", True)
    line(3, "Đất thừa", "Vthua", f"={vd}-{vdp}", "Vdao − Vdap", "m3", True)
    line("", "Chiều sâu hố đào lớn nhất", "Hmax", f"=MAX({','.join(hd.values())})", "max(Hd)", "m")
    line("", "Cạnh đền bù", "Dbu",
         "=T1_Bt+MAX(" + ",".join(f"T{x['n']}_b2-T{x['n']}_b1+T{x['n']}_Bm" for x in info) + ")+2*0.5",
         "Bt + max(b2 − b1 + Bm) + 2·0.5", "m")
    r += 2
    c = ws.cell(row=r, column=2, value=chk)
    c.font = Font(bold=True, color="C00000")
    return out


def _summary(wb, mp, info, dat):
    """Sheet TONG KE: tổng kê cả vị trí móng (cộng 2 loại)."""
    ws = wb.create_sheet("TONG KE", 0)
    D = "DAT!"
    kl = [f"'{x['skl']}'!" for x in info]
    used = sorted({int(d + 0.5) for x in info for d, _, w in x["g"].klc if w > 1e-6})
    cols = [(f"{mp['BTM']} (m³)", "=" + "+".join(f"{k}$E${x['rows']['Vbt']}" for k, x in zip(kl, info)), "0.00"),
            (f"{mp['BTL']} (m³)", "=" + "+".join(f"{k}$E${x['rows']['Vlot']}" for k, x in zip(kl, info)), "0.00")]
    for d in used:
        cols.append((f"Ø{d} (kg)", "=" + "+".join(f"{k}$H${x['rd'][d]}" for k, x in zip(kl, info) if d in x["rd"]),
                     "0.00"))
    cols += [("TỔNG K.LƯỢNG THÉP (kg)", "=" + "+".join(f"{k}$H${x['rtot']}" for k, x in zip(kl, info)), "0.00"),
             ("BULONG NEO", lm.bolt_text(mp), None),
             ("Đào đất (m³)", f"={D}{dat['Vdao']}", "0.00"),
             ("Đắp đất (m³)", f"={D}{dat['Vdap']}", "0.00"),
             ("Ván khuôn (m²)", "=" + "+".join(f"{k}$E${x['rows']['VK']}" for k, x in zip(kl, info)), "0.00"),
             ("Rộng đền bù (m)", f"={D}{dat['Dbu']}", "0.00"),
             ("Dài đền bù (m)", f"={D}{dat['Dbu']}", "0.00"),
             ("Diện đền bù (m²)", None, "0.00"),
             ("Chiều sâu hố đào (m)", f"={D}{dat['Hmax']}", "0.00")]
    n = len(cols)
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 64
    ws.row_dimensions[3].height = 24
    c = ws.cell(row=1, column=1, value="BẢNG TỔNG KÊ KHỐI LƯỢNG CHO 01 VỊ TRÍ MÓNG")
    c.font = Font(size=14)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    for j in range(1, n + 1):
        ws.cell(row=1, column=j).border = K.BOX
    for j, (h, f, fmt) in enumerate(cols, 1):
        ws.column_dimensions[L(j)].width = 11 if len(h) < 14 else 12.5
        K._cell(ws, 2, j, h, align=Alignment(horizontal="center", vertical="center", wrap_text=True))
        if h.startswith("Diện đền bù"):
            f = f"={L(j - 2)}3*{L(j - 1)}3"
        K._cell(ws, 3, j, f, fmt, align=Alignment(horizontal="center", vertical="center"))
    names = ", ".join(f"loại {x['n']}: {len(mp['TYPES'][x['n'] - 1]['POS'])}{mp['TYPES'][x['n'] - 1]['NAME']}"
                      for x in info)
    ws.cell(row=5, column=1, value=f"Cả vị trí móng ({names}). Khối lượng thép đã gồm dự phòng.").font = \
        Font(italic=True, color="808080")
    jd = [j for j, (h, _, _) in enumerate(cols, 1) if h.startswith("Ø")]
    jt = next(j for j, (h, _, _) in enumerate(cols, 1) if h.startswith("TỔNG K.LƯỢNG THÉP"))
    sm = "+".join(f"{L(j)}3" for j in jd) or "0"
    c = ws.cell(row=6, column=1, value=f'=IF(ABS(({sm})-{L(jt)}3)>0.01,"Có đường kính mới - xuất lại file","")')
    c.font = Font(bold=True, color="C00000")
    c = ws.cell(row=7, column=1, value="=DAT!B" + str(_chk_row(wb)))
    c.font = Font(bold=True, color="C00000")


def _chk_row(wb):
    ws = wb["DAT"]
    for r in range(ws.max_row, 0, -1):
        v = ws.cell(r, 2).value
        if isinstance(v, str) and v.startswith("=IF(MIN("):
            return r
    return ws.max_row
