"""Đối chiếu Excel khối lượng móng 2 loại (công thức, LibreOffice tính lại) với chương trình.

    python tools/check_kl_lech.py [thư_mục_tạm]
"""
import os
import subprocess
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
import openpyxl  # noqa: E402

from mongtru.lech import export as le  # noqa: E402
from mongtru.lech import kl_excel as lk  # noqa: E402
from mongtru.lech import model as lm  # noqa: E402

OUT = sys.argv[1] if len(sys.argv) > 1 else tempfile.mkdtemp()
os.makedirs(OUT, exist_ok=True)
SAMPLE = os.path.join(ROOT, "samples", "N112-52(60) - 3T-1T.xls")
# (thông số chung, [thay đổi loại 1, thay đổi loại 2])
CASES = {
    "mau": ({}, [{}, {}]),
    "cao_do": ({}, [dict(CD=-2.0), {}]),
    "2_2": ({}, [dict(POS=[1, 2], B=3.6, CD=-1.0), dict(POS=[3, 4])]),
    "chung_bac": (dict(BT=5.5), [dict(CD=-1.5, B2=1.6), dict(B=3.6)]),
    "chung_bang": (dict(BT=5.0), [{}, dict(B=3.8)]),
    "vat_ca2": (dict(DP=5.0), [dict(B2=1.5, RC_N=28), dict(B2=1.3, RB_BOT_S=150)]),
}
base, _ = lm.read_excel(SAMPLE)
mps = {}
for nm, (ov, tov) in CASES.items():
    mp = {k: (v if k != "TYPES" else [dict(t) for t in v]) for k, v in base.items()}
    mp.update(ov)
    for t, o in zip(mp["TYPES"], tov):
        t.update(o)
    assert lm.validate(mp) is None, (nm, lm.validate(mp))
    gs = le.geos(mp)
    lk.export(mp, gs, os.path.join(OUT, f"{nm}.xlsx"))
    mps[nm] = (mp, gs)
calc = os.path.join(OUT, "calc")
subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", calc] +
               [os.path.join(OUT, f"{nm}.xlsx") for nm in CASES], capture_output=True, timeout=900)
bad = 0
for nm, (mp, gs) in mps.items():
    wb = openpyxl.load_workbook(os.path.join(calc, f"{nm}.xlsx"), data_only=True)
    # thép từng loại
    for i, g in enumerate(gs):
        ws = wb[f"KHOI LUONG {i + 1}"]
        rows = {}
        for r in range(1, ws.max_row + 1):
            b = ws.cell(r, 2).value
            if isinstance(b, str) and ws.cell(r, 9).value and isinstance(ws.cell(r, 5).value, (int, float)):
                rows[b.split(" - ")[0]] = [ws.cell(r, c).value for c in (3, 4, 5, 8)]
        for r in g.kl:
            ex = rows.get(r["name"])
            want = [int(r["dia"] + 0.5), r["len"], r["n"], r["kg"]]
            ok = ex is not None and ex[0] == want[0] and abs(ex[1] - want[1]) <= 0.6 and ex[2] == want[2] and \
                abs(ex[3] - want[3] * round(r["len"]) / r["len"]) <= max(0.05, 0.0005 * want[3])
            if not ok:
                bad += 1
                print(f"[{nm}] loại {i + 1} vị trí {r['pos']} {r['name']}: Excel {ex} - chương trình {want}")
    tk = wb["TONG KE"]
    vals = {tk.cell(2, j).value: tk.cell(3, j).value for j in range(1, tk.max_column + 1)}
    dpk = 1 + mp["DP"] / 100
    steel = sum(r["kg"] for g in gs for r in g.kl) * dpk
    want = {f"{mp['BTM']} (m³)": lm.vol_mong(mp), f"{mp['BTL']} (m³)": lm.vol_lot(mp),
            "Đào đất (m³)": lm.vol_dao(mp), "Đắp đất (m³)": lm.vol_dap(mp), "Ván khuôn (m²)": lm.area_vk(mp),
            "TỔNG K.LƯỢNG THÉP (kg)": steel, "Rộng đền bù (m)": lm.den_bu(mp),
            "Chiều sâu hố đào (m)": lm.pit_depth_max(mp)}
    for key, v in want.items():
        if abs((vals.get(key) or 1e9) - v) > max(0.01, 0.0005 * v):
            bad += 1
            print(f"[{nm}] {key}: Excel {vals.get(key)} - chương trình {v:.3f}")
    for msg in (tk.cell(6, 1).value, tk.cell(7, 1).value):
        if msg:
            bad += 1
            print(f"[{nm}] TONG KE báo: {msg}")
    print(f"[{nm}] đào {'chung' if lm.pit_joint(mp) else 'riêng'}: {lm.vol_dao(mp):.3f} m3")
print("Số lỗi:", bad, "/", len(CASES), "bộ thông số")
sys.exit(1 if bad else 0)
